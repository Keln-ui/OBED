import io
import os
import tempfile
from unittest.mock import patch

from docx import Document
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse
from PIL import Image
from PyPDF2 import PdfReader
from reportlab.pdfgen import canvas


class DocumentConversionTests(TestCase):
    def setUp(self):
        self.media_dir = tempfile.TemporaryDirectory()
        self.media_settings = patch('django.conf.settings.MEDIA_ROOT', self.media_dir.name)
        self.media_settings.start()
        self.client = Client()

    def tearDown(self):
        self.media_settings.stop()
        self.media_dir.cleanup()

    @staticmethod
    def make_pdf():
        output = io.BytesIO()
        pdf = canvas.Canvas(output)
        pdf.setFont('Helvetica-Bold', 18)
        pdf.drawString(72, 740, 'Layout preservation test')
        pdf.setFont('Helvetica', 12)
        pdf.drawString(72, 710, 'This source content should remain editable in the Word document.')
        pdf.save()
        return output.getvalue()

    def test_pdf_upload_returns_a_valid_editable_word_document(self):
        pdf_file = SimpleUploadedFile(
            'layout-test.pdf',
            self.make_pdf(),
            content_type='application/pdf',
        )
        with (
            patch('pdf_converter.views.get_conversion_cost', return_value=2),
            patch('pdf_converter.views.payment_required_for_conversion', return_value=False),
        ):
            response = self.client.post(
                reverse('upload_pdf'),
                {'pdf_file': pdf_file, 'output_format': 'docx'},
            )

        self.assertEqual(response.status_code, 200, response.content)
        result = response.json()
        self.assertTrue(result['success'], result)
        self.assertTrue(result['filename'].endswith('.docx'))
        output_path = os.path.join(self.media_dir.name, 'converted', result['filename'])
        self.assertTrue(os.path.isfile(output_path))
        document = Document(output_path)
        content = '\n'.join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn('Layout preservation test', content)
        self.assertIn('remain editable', content)

    def test_unsupported_word_97_format_is_rejected(self):
        response = self.client.post(
            reverse('upload_pdf'),
            {
                'pdf_file': SimpleUploadedFile(
                    'layout-test.pdf',
                    self.make_pdf(),
                    content_type='application/pdf',
                ),
                'output_format': 'doc',
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('supported output format', response.json()['error'])

    def test_word_to_pdf_download_contains_source_document_text(self):
        source = Document()
        source.add_heading('Original document heading', level=1)
        source.add_paragraph('The source paragraph must appear in the PDF output.')
        docx_bytes = io.BytesIO()
        source.save(docx_bytes)
        upload = SimpleUploadedFile(
            'source.docx',
            docx_bytes.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )

        with patch('pdf_converter.converters.docx_to_pdf_converter.shutil.which', return_value=None):
            response = self.client.post(reverse('upload_docx_to_pdf'), {'docx_file': upload})

        self.assertEqual(response.status_code, 200, response.content)
        result = response.json()
        self.assertTrue(result['success'], result)
        download_name = result['download_url'].rstrip('/').rsplit('/', 1)[-1]
        self.assertEqual(result['filename'], download_name)
        download = self.client.get(result['download_url'])
        self.assertEqual(download.status_code, 200)
        self.assertTrue(download['Content-Disposition'].endswith(result['filename'] + '"'))
        self.assertEqual(download['Content-Type'], 'application/pdf')
        try:
            output_data = b''.join(download.streaming_content)
        finally:
            download.close()
        output_text = '\n'.join(
            page.extract_text() or ''
            for page in PdfReader(io.BytesIO(output_data)).pages
        )
        self.assertIn('Original document heading', output_text)
        self.assertIn('source paragraph', output_text)

    def test_word_to_pdf_response_filename_matches_download_path(self):
        source = Document()
        source.add_paragraph('Filename contract')
        docx_bytes = io.BytesIO()
        source.save(docx_bytes)
        upload = SimpleUploadedFile('filename.docx', docx_bytes.getvalue())

        with patch('pdf_converter.converters.docx_to_pdf_converter.shutil.which', return_value=None):
            response = self.client.post(reverse('upload_docx_to_pdf'), {'docx_file': upload})

        self.assertEqual(response.status_code, 200, response.content)
        result = response.json()
        self.assertEqual(
            result['filename'],
            result['download_url'].rstrip('/').rsplit('/', 1)[-1],
        )
        self.assertTrue(result['filename'].startswith('docx2pdf_'))

    def test_image_to_pdf_button_has_a_valid_download(self):
        image_data = io.BytesIO()
        Image.new('RGB', (20, 20), color='blue').save(image_data, format='PNG')
        upload = SimpleUploadedFile('picture.png', image_data.getvalue(), content_type='image/png')

        response = self.client.post(reverse('upload_image_to_pdf'), {'image_file': upload})
        self.assertEqual(response.status_code, 200, response.content)
        result = response.json()
        self.assertTrue(result['success'], result)
        self.assertEqual(result['filename'], result['download_url'].rstrip('/').rsplit('/', 1)[-1])

        download = self.client.get(result['download_url'])
        self.assertEqual(download.status_code, 200)
        self.assertEqual(download['Content-Type'], 'application/pdf')
        self.assertTrue(download['Content-Disposition'].endswith(result['filename'] + '"'))
        try:
            output_data = b''.join(download.streaming_content)
        finally:
            download.close()
        self.assertEqual(len(PdfReader(io.BytesIO(output_data)).pages), 1)

    def test_word_to_pdf_splits_oversized_table_rows_across_pages(self):
        source = Document()
        table = source.add_table(rows=1, cols=3)
        table.cell(0, 0).text = 'Long cell text. ' * 500
        table.cell(0, 1).text = 'Symbols A & B should be preserved.'
        table.cell(0, 2).text = 'End of row.'
        docx_bytes = io.BytesIO()
        source.save(docx_bytes)
        upload = SimpleUploadedFile(
            'long-table.docx',
            docx_bytes.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )

        with patch('pdf_converter.converters.docx_to_pdf_converter.shutil.which', return_value=None):
            response = self.client.post(reverse('upload_docx_to_pdf'), {'docx_file': upload})

        self.assertEqual(response.status_code, 200, response.content)
        result = response.json()
        download = self.client.get(result['download_url'])
        try:
            output_data = b''.join(download.streaming_content)
        finally:
            download.close()
        output = PdfReader(io.BytesIO(output_data))
        self.assertGreater(len(output.pages), 1)
        output_text = '\n'.join(page.extract_text() or '' for page in output.pages)
        self.assertIn('Symbols A & B should be preserved.', output_text)
