import os
import logging
from .base_converter import BaseConverter

logger = logging.getLogger(__name__)


class PPTConverter(BaseConverter):
    """Convert PDF to PowerPoint presentation using PyMuPDF and python-pptx"""
    
    def convert(self):
        """Convert PDF to PPTX by rendering each page as an image on a slide"""
        try:
            import fitz  # PyMuPDF
            from pptx import Presentation
            from pptx.util import Inches
            
            # Open the PDF
            pdf_document = fitz.open(str(self.pdf_path))
            
            # Create a new PowerPoint presentation
            prs = Presentation()
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)
            
            for page_num in range(len(pdf_document)):
                page = pdf_document[page_num]
                
                # Add a blank slide
                blank_layout = prs.slide_layouts[6]  # Blank layout
                slide = prs.slides.add_slide(blank_layout)
                
                # Render the page as an image at high resolution
                # Use a higher zoom factor for better quality
                zoom = 2
                mat = fitz.Matrix(zoom, zoom)
                pix = page.get_pixmap(matrix=mat)
                
                # Save the rendered page as a temporary image
                temp_img_path = f"{self.output_path}.temp_page_{page_num}.png"
                pix.save(temp_img_path)
                
                # Add the image to the slide (fit to slide dimensions)
                left = Inches(0.5)
                top = Inches(0.5)
                width = Inches(9)
                height = Inches(6.5)
                
                try:
                    slide.shapes.add_picture(temp_img_path, left, top, width, height)
                except Exception as e:
                    logger.warning(f"Could not add image to slide: {str(e)}")
                finally:
                    # Clean up temp image
                    if os.path.exists(temp_img_path):
                        os.remove(temp_img_path)
            
            pdf_document.close()
            
            # Create output directory if it doesn't exist
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            
            # Save the presentation
            prs.save(str(self.output_path))
            
            logger.info(f"Successfully converted PDF to PPTX: {self.output_path}")
            return True
            
        except ImportError as e:
            logger.error(f"Missing required library for PPT conversion: {str(e)}")
            self.error_message = f"Missing library: {str(e)}. Please install python-pptx and PyMuPDF."
            return False
        except Exception as e:
            logger.error(f"Error converting PDF to PPTX: {str(e)}")
            self.error_message = str(e)
            return False
