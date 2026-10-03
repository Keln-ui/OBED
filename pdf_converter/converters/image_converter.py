import os
import logging
import tempfile
from .base_converter import BaseConverter

logger = logging.getLogger(__name__)


class ImageConverter(BaseConverter):
    """Converter for image files (PNG, JPG, etc.) directly to Word using OCR"""
    
    def convert(self) -> bool:
        """Convert image to Word document using OCR"""
        try:
            # Create output directory if it doesn't exist
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            
            # Try to use pytesseract for direct OCR (best quality)
            try:
                import pytesseract
                from PIL import Image
                from docx import Document
                from docx.shared import Inches, Pt
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                
                # Open the image
                image = Image.open(self.pdf_path)
                
                # Perform OCR to extract text
                ocr_text = pytesseract.image_to_string(image)
                
                # Create a new Word document
                doc = Document()
                
                # Add a title with the original filename
                title = doc.add_heading(os.path.splitext(os.path.basename(self.pdf_path))[0], 0)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                # Add the OCR text as paragraphs
                if ocr_text.strip():
                    paragraphs = ocr_text.strip().split('\n')
                    for para_text in paragraphs:
                        if para_text.strip():
                            doc.add_paragraph(para_text)
                else:
                    doc.add_paragraph("[No text detected in image]")
                
                # Save the document
                doc.save(str(self.output_path))
                
                logger.info(f"Successfully converted image {self.pdf_path} to {self.output_path} using pytesseract")
                return True
                
            except ImportError as e:
                missing_pkg = str(e).split("'")[1] if "'" in str(e) else "required package"
                logger.warning(f"Direct OCR unavailable ({missing_pkg} not installed), trying ocrmypdf + text extraction fallback")
                return self._convert_via_ocr_pdf()
                
        except Exception as e:
            logger.error(f"Error converting image to Word: {str(e)}")
            return False
    
    def _convert_via_ocr_pdf(self) -> bool:
        """Fallback: Convert image to PDF, OCR with ocrmypdf, extract text, then create Word doc"""
        temp_pdf_path = None
        ocr_pdf_path = None
        
        try:
            from PIL import Image
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            import PyPDF2
            
            # Create temporary PDF from image
            with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as temp_pdf:
                temp_pdf_path = temp_pdf.name
            
            # Convert image to PDF
            with Image.open(self.pdf_path) as img:
                img.convert('RGB').save(temp_pdf_path, 'PDF')
            
            # Try to OCR the PDF using ocrmypdf
            try:
                import subprocess
                
                with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as ocr_pdf:
                    ocr_pdf_path = ocr_pdf.name
                
                result = subprocess.run(
                    ['ocrmypdf', '--output-type', 'pdf', '--force-ocr', temp_pdf_path, ocr_pdf_path],
                    capture_output=True,
                    text=True,
                    timeout=300
                )
                
                if result.returncode == 0:
                    # OCR successful, extract text from OCR'd PDF
                    text_content = []
                    with open(ocr_pdf_path, 'rb') as file:
                        pdf_reader = PyPDF2.PdfReader(file)
                        for page in pdf_reader.pages:
                            text = page.extract_text()
                            if text:
                                text_content.append(text)
                    
                    # Create Word document with extracted text
                    doc = Document()
                    title = doc.add_heading(os.path.splitext(os.path.basename(self.pdf_path))[0], 0)
                    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    if text_content:
                        for page_text in text_content:
                            paragraphs = page_text.strip().split('\n')
                            for para_text in paragraphs:
                                if para_text.strip():
                                    doc.add_paragraph(para_text)
                    else:
                        doc.add_paragraph("[No text could be extracted from the image]")
                    
                    doc.save(str(self.output_path))
                    logger.info(f"Successfully converted image {self.pdf_path} to {self.output_path} via OCR PDF + text extraction")
                    return True
                else:
                    logger.warning(f"ocrmypdf failed: {result.stderr}")
                    raise RuntimeError("OCR failed")
                    
            except (FileNotFoundError, RuntimeError):
                logger.warning("OCR unavailable, creating Word document with image embedded")
                return self._convert_image_to_word_with_image()
                    
        except Exception as e:
            logger.error(f"Error in fallback image conversion: {str(e)}")
            return False
        finally:
            # Clean up temporary files
            for path in [temp_pdf_path, ocr_pdf_path]:
                if path and os.path.exists(path):
                    try:
                        os.remove(path)
                    except Exception as e:
                        logger.warning(f"Failed to remove temporary file {path}: {e}")
    
    def _convert_image_to_word_with_image(self) -> bool:
        """Last resort: Create Word document with the image embedded (no OCR text)"""
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            
            doc = Document()
            title = doc.add_heading(os.path.splitext(os.path.basename(self.pdf_path))[0], 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            doc.add_paragraph("[This image was converted without OCR. Install pytesseract or ocrmypdf for text extraction.]")
            doc.add_paragraph("")
            
            # Add the image to the document
            doc.add_picture(self.pdf_path, width=Inches(6))
            
            doc.save(str(self.output_path))
            logger.info(f"Created Word document with embedded image {self.pdf_path}")
            return True
            
        except Exception as e:
            logger.error(f"Error creating Word document with image: {str(e)}")
            return False
