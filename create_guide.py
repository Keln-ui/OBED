#!/usr/bin/env python
"""
Script to create Core Synch Solutions System Guide as a proper Word document
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

def set_cell_border(cell, **kwargs):
    """Helper function to set cell borders"""
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            edge_el = tcPr.find(qn(f'w:{edge}'))
            if edge_el is None:
                edge_el = docx.oxml.OxmlElement(f'w:{edge}')
                tcPr.append(edge_el)
            edge_el.set(qn('w:val'), kwargs[edge])
            edge_el.set(qn('w:sz'), '4')
            edge_el.set(qn('w:space'), '0')
            edge_el.set(qn('w:color'), 'auto')

def add_heading_custom(doc, text, level=1):
    """Add a styled heading"""
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        run.font.color.rgb = RGBColor(102, 126, 234)  # Purple color
        run.font.bold = True
    return heading

def add_bullet_point(doc, text, indent_level=0):
    """Add a bullet point"""
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
    p.paragraph_format.first_line_indent = Inches(-0.25)
    return p

# Create document
doc = Document()

# Set default font
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

# Title Page
title = doc.add_heading('Core Synch Solutions', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.color.rgb = RGBColor(102, 126, 234)
    run.font.size = Pt(32)
    run.font.bold = True

subtitle = doc.add_paragraph('Complete System Guide')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(18)
subtitle.runs[0].font.color.rgb = RGBColor(102, 126, 234)
subtitle.runs[0].font.bold = True

doc.add_paragraph()
version = doc.add_paragraph('Version 1.0')
version.alignment = WD_ALIGN_PARAGRAPH.CENTER
version.runs[0].font.color.rgb = RGBColor(100, 100, 100)

date = doc.add_paragraph('June 2026')
date.alignment = WD_ALIGN_PARAGRAPH.CENTER
date.runs[0].font.color.rgb = RGBColor(100, 100, 100)

doc.add_page_break()

# Table of Contents
add_heading_custom(doc, 'Table of Contents', level=1)
toc_items = [
    '1. Introduction',
    '2. Getting Started',
    '3. Document Conversion',
    '4. Revision Materials',
    '5. Subscriptions & Plans',
    '6. User Dashboard',
    '7. Admin Features',
    '8. FAQs'
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.left_indent = Inches(0.25)

doc.add_page_break()

# 1. Introduction
add_heading_custom(doc, '1. Introduction', level=1)
doc.add_paragraph('Core Synch Solutions is a comprehensive educational platform that provides:')
add_bullet_point(doc, 'Document Conversion: Convert PDFs and images to various formats (Word, PowerPoint, HTML, etc.)')
add_bullet_point(doc, 'Revision Materials: Access past papers, marking schemes, notes, and textbooks')
add_bullet_point(doc, 'Flexible Subscriptions: Individual and school-based plans with different feature access')

doc.add_paragraph()
p = doc.add_paragraph('Key Features:')
p.runs[0].font.bold = True
p.runs[0].font.color.rgb = RGBColor(102, 126, 234)

features = [
    'Convert PDF to DOCX, DOC, TXT, HTML, JSON, CSV, PPTX',
    'Free Word-to-PDF conversion',
    'OCR for scanned documents',
    'Revision materials for CBE and 8-4-4 curricula',
    'Individual and school subscription plans',
    'Advanced search and filtering',
    'User dashboard with activity tracking'
]
for feature in features:
    add_bullet_point(doc, feature)

# 2. Getting Started
add_heading_custom(doc, '2. Getting Started', level=1)
add_heading_custom(doc, 'For Individual Users', level=2)

doc.add_paragraph('Step 1: Create an Account')
steps = [
    'Visit the homepage: http://127.0.0.1:8000/',
    'Click "Sign Up" in the top navigation',
    'Fill in your details: Username, Email address, Password (minimum 8 characters)',
    'Click "Create Account"'
]
for step in steps:
    add_bullet_point(doc, step)

doc.add_paragraph('Step 2: Browse Features')
add_bullet_point(doc, 'Document Converter: Upload and convert PDFs')
add_bullet_point(doc, 'Revision Materials: Browse free materials (no subscription required)')
add_bullet_point(doc, 'Subscription Plans: Upgrade for premium features')

add_heading_custom(doc, 'For School Administrators', level=2)
doc.add_paragraph('Step 1: Register Your School')
steps = [
    'Visit the homepage',
    'Click "Sign Up"',
    'Select "Register a School instead" link',
    'Fill in school details: School name, address, phone, email, principal name',
    'Create username and password',
    'Click "Create Account"'
]
for step in steps:
    add_bullet_point(doc, step)

doc.add_paragraph('Step 2: Subscribe to a Plan')
add_bullet_point(doc, 'After registration, you\'ll be redirected to subscription plans')
add_bullet_point(doc, 'Choose a plan suitable for your school')
add_bullet_point(doc, 'Complete payment via M-Pesa')
add_bullet_point(doc, 'Subscription activates immediately')

# 3. Document Conversion
add_heading_custom(doc, '3. Document Conversion', level=1)
add_heading_custom(doc, 'Free Conversion: Word to PDF', level=2)

p = doc.add_paragraph('No payment required!')
p.runs[0].font.bold = True
p.runs[0].font.color.rgb = RGBColor(16, 185, 129)

steps = [
    'Scroll to the "Word to PDF Converter" section (green card)',
    'Click "Choose File" or drag-and-drop a .docx file',
    'Click "Convert to PDF"',
    'Wait for conversion to complete',
    'Click "Download PDF"'
]
for step in steps:
    add_bullet_point(doc, step)

p = doc.add_paragraph('Supported format: .docx only')
p.runs[0].font.italic = True

add_heading_custom(doc, 'Paid Conversion: PDF to Other Formats', level=2)
p = doc.add_paragraph('Cost: KSH 2 per page')
p.runs[0].font.bold = True

doc.add_paragraph('Step 1: Upload File')
add_bullet_point(doc, 'Go to the main converter section')
add_bullet_point(doc, 'Click "Choose File" or drag-and-drop a PDF/image')
add_bullet_point(doc, 'Supported formats: PDF, PNG, JPG, JPEG, GIF, BMP')

doc.add_paragraph('Step 2: Configure Settings')
add_bullet_point(doc, 'Output Format: Select from Word (.docx), Word 97-2003 (.doc), Text (.txt), HTML, JSON, CSV, PowerPoint (.pptx)')
add_bullet_point(doc, 'Use OCR: Check this box for scanned documents (requires subscription)')

doc.add_paragraph('Step 3: Convert and Pay')
add_bullet_point(doc, 'Click "Convert Now"')
add_bullet_point(doc, 'System shows estimated pages and cost')
add_bullet_point(doc, 'After conversion, click "Pay with M-Pesa"')
add_bullet_point(doc, 'Enter your M-Pesa phone number')
add_bullet_point(doc, 'Complete payment')
add_bullet_point(doc, 'Download your converted file')

# 4. Revision Materials
add_heading_custom(doc, '4. Revision Materials', level=1)
add_heading_custom(doc, 'Browsing Materials', level=2)

doc.add_paragraph('Option 1: Browse by Curriculum')
add_bullet_point(doc, 'Click "Revision" in the navigation menu')
add_bullet_point(doc, 'Choose your curriculum: CBE (Competency-Based Education) or 8-4-4 (Traditional)')
add_bullet_point(doc, 'Select your level: Primary, Secondary, or College/University')
add_bullet_point(doc, 'Click on a class to view subjects')
add_bullet_point(doc, 'Select a subject to see available materials')

doc.add_paragraph('Option 2: Browse by Subject')
add_bullet_point(doc, 'From Revision home, click "Browse by Subject"')
add_bullet_point(doc, 'See all subjects with available materials')
add_bullet_point(doc, 'Filter by curriculum (CBE or 8-4-4)')
add_bullet_point(doc, 'Click on a subject to view materials')

doc.add_paragraph('Option 3: Advanced Search')
add_bullet_point(doc, 'Click "Search" in the navigation menu')
add_bullet_point(doc, 'Use filters: Keyword, Type, Class, Term, Year, Price, Sort')
add_bullet_point(doc, 'Click "Search" to see results')

add_heading_custom(doc, 'Material Types', level=2)
add_bullet_point(doc, 'Past Papers: Previous exam papers')
add_bullet_point(doc, 'Marking Schemes: Official answers and grading rubrics')
add_bullet_point(doc, 'Notes: Study notes and summaries')
add_bullet_point(doc, 'Textbooks: Digital textbook resources')

add_heading_custom(doc, 'Accessing Materials', level=2)
p = doc.add_paragraph('Free Materials')
p.runs[0].font.bold = True
add_bullet_point(doc, 'No subscription required')
add_bullet_point(doc, 'Available to all users (including anonymous)')
add_bullet_point(doc, 'Click "View" to see details')
add_bullet_point(doc, 'Click "Download" to get the file')

p = doc.add_paragraph('Premium Materials')
p.runs[0].font.bold = True
add_bullet_point(doc, 'Requires active subscription with revision access')
add_bullet_point(doc, 'Click "View" to see details')
add_bullet_point(doc, 'If you have access: Download immediately')
add_bullet_point(doc, 'If not: You\'ll be prompted to subscribe')

# 5. Subscriptions & Plans
add_heading_custom(doc, '5. Subscriptions & Plans', level=1)
add_heading_custom(doc, 'Plan Types', level=2)

p = doc.add_paragraph('For Individual Users')
p.runs[0].font.bold = True
add_bullet_point(doc, 'Personal Learning Plan: KSH 500/year - Revision materials + PDF conversions')
add_bullet_point(doc, 'Student Premium: KSH 1,000/year - Full access including OCR')

p = doc.add_paragraph('For Schools')
p.runs[0].font.bold = True
add_bullet_point(doc, 'School Basic: KSH 2,000/year - PDF conversions only, up to 10 teachers')
add_bullet_point(doc, 'School Standard: KSH 3,500/year - Conversions + Revision, unlimited teachers')
add_bullet_point(doc, 'School Premium: KSH 5,000/year - Full access including OCR')

add_heading_custom(doc, 'Subscribing to a Plan', level=2)
doc.add_paragraph('Step 1: View Available Plans')
add_bullet_point(doc, 'Click "Subscription Plans" in navigation')
add_bullet_point(doc, 'See plans available for your account type')

doc.add_paragraph('Step 2: Select a Plan')
add_bullet_point(doc, 'Click "Subscribe Now" on your chosen plan')
add_bullet_point(doc, 'Review plan details: Name, price, duration, features')
add_bullet_point(doc, 'Click "Subscribe"')

doc.add_paragraph('Step 3: Complete Payment')
add_bullet_point(doc, 'Enter your M-Pesa phone number')
add_bullet_point(doc, 'Click "Pay with M-Pesa"')
add_bullet_point(doc, 'You\'ll receive an M-Pesa prompt on your phone')
add_bullet_point(doc, 'Enter your PIN to complete payment')
add_bullet_point(doc, 'Subscription activates immediately')

# 6. User Dashboard
add_heading_custom(doc, '6. User Dashboard', level=1)
doc.add_paragraph('Access your dashboard by clicking "Dashboard" in the navigation menu (when logged in).')

add_heading_custom(doc, 'Dashboard Sections', level=2)
add_bullet_point(doc, 'Stats Overview: Total conversions, total spent, materials downloaded, subscription status')
add_bullet_point(doc, 'Recent Conversions: Last 10 conversions with download options')
add_bullet_point(doc, 'Quick Actions: New conversion, revision materials, subscription, profile')
add_bullet_point(doc, 'Popular Revision Materials: Trending free materials')

# 7. Admin Features
add_heading_custom(doc, '7. Admin Features', level=1)
doc.add_paragraph('Access admin dashboard by logging in with staff/admin credentials and clicking "Dashboard" in navigation.')

add_heading_custom(doc, 'Admin Dashboard Sections', level=2)
add_bullet_point(doc, 'Statistics Overview: Total conversions, revenue, materials, downloads')
add_bullet_point(doc, 'Weekly Performance: Daily activity and revenue')
add_bullet_point(doc, 'Recent Activity: Latest conversions and payments')
add_bullet_point(doc, 'Attention Needed: Pending payments and issues')
add_bullet_point(doc, 'Quick Actions: Manage plans, upload materials, add subjects')

add_heading_custom(doc, 'Django Admin Panel', level=2)
doc.add_paragraph('Access at: http://127.0.0.1:8000/admin/')
add_bullet_point(doc, 'Manage Subscription Plans: Create plans with different targets (Individual/School/Both)')
add_bullet_point(doc, 'Upload Revision Materials: Add past papers, notes, textbooks')
add_bullet_point(doc, 'Link Materials: Connect past papers to marking schemes')
add_bullet_point(doc, 'Manage Users and Schools: View and edit user accounts')

# 8. FAQs
add_heading_custom(doc, '8. FAQs', level=1)

add_heading_custom(doc, 'General Questions', level=2)
faqs = [
    ('Q: Do I need an account?', 'A: No, you can browse free materials without an account. Account needed for conversions and premium materials.'),
    ('Q: Individual vs School subscription?', 'A: Individual is for personal use. School is for entire institution with multiple teachers.'),
    ('Q: Can I upgrade later?', 'A: Yes! Go to Subscription Plans and select a higher tier anytime.')
]
for q, a in faqs:
    p = doc.add_paragraph()
    p.add_run(q).font.bold = True
    doc.add_paragraph(a)

add_heading_custom(doc, 'Document Conversion', level=2)
faqs = [
    ('Q: How much does conversion cost?', 'A: KSH 2 per page. Word-to-PDF is FREE.'),
    ('Q: What is OCR?', 'A: Optical Character Recognition for scanned documents. Requires premium subscription.'),
    ('Q: How long does conversion take?', 'A: Usually 10-30 seconds depending on file size.')
]
for q, a in faqs:
    p = doc.add_paragraph()
    p.add_run(q).font.bold = True
    doc.add_paragraph(a)

add_heading_custom(doc, 'Revision Materials', level=2)
faqs = [
    ('Q: Are all materials free?', 'A: No, some require subscription with revision access.'),
    ('Q: How to identify free vs paid?', 'A: Green "Free" badge = no subscription needed. Yellow "KSH X" = requires subscription.'),
    ('Q: Can I preview materials?', 'A: Yes, click "View" to see preview with first pages.')
]
for q, a in faqs:
    p = doc.add_paragraph()
    p.add_run(q).font.bold = True
    doc.add_paragraph(a)

add_heading_custom(doc, 'Subscriptions', level=2)
faqs = [
    ('Q: How do I pay?', 'A: We use M-Pesa. Enter your phone number during checkout.'),
    ('Q: How long does subscription last?', 'A: Termly = 3 months, Yearly = 12 months.'),
    ('Q: What happens when it expires?', 'A: You lose premium access. Data and conversions are preserved.')
]
for q, a in faqs:
    p = doc.add_paragraph()
    p.add_run(q).font.bold = True
    doc.add_paragraph(a)

# Contact Information
doc.add_page_break()
add_heading_custom(doc, 'Contact Information', level=1)
p = doc.add_paragraph('Company: Core Synch Solutions')
p.runs[0].font.bold = True
doc.add_paragraph('Platform: Educational Document Management System')
doc.add_paragraph('Version: 1.0')
doc.add_paragraph('Last Updated: June 2026')

# Save document
doc.save('CORESYNCH_SYSTEM_GUIDE.docx')
print("Document created successfully: CORESYNCH_SYSTEM_GUIDE.docx")